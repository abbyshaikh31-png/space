import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rows = open('addresses.csv').read().strip().split('\n')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Addresses'
for line in rows:
    ws.append(line.split(','))
for c in range(1, 4):
    cell = ws.cell(1, c)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='4472C4')
    cell.alignment = Alignment(horizontal='center')
for c in range(1, 4):
    ws.column_dimensions[get_column_letter(c)].width = 36
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:C%d' % len(rows)
wb.save('addresses.xlsx')
print('rows:', len(rows)-1, 'size:', os.path.getsize('addresses.xlsx'))
