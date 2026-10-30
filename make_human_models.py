"""Synthetic human profile generator -> Excel.
Generates realistic-looking (fully synthetic) person records. NO real-person data.
Run: python3 make_human_models.py  (then open human_models.xlsx)
"""
import csv, random
from datetime import date, timedelta

random.seed(42)
N = 500  # number of synthetic human records

FIRST_NAMES = ["Arjun","Priya","Rahul","Sneha","Vikram","Ananya","Rohan","Kavya","Rajesh","Meera",
    "Neha","Amit","Divya","Kiran","Pooja","Sanjay","Lakshmi","Deepak","Anita","Farhan",
    "Zara","Liam","Emma","Noah","Olivia","Ethan","Ava","Lucas","Mia","Aiden"]
LAST_NAMES = ["Sharma","Patel","Iyer","Reddy","Nair","Gupta","Singh","Khan","Mehta","Joshi",
    "Das","Rao","Kulkarni","Pillai","Bose","Chopra","Malhotra","Verma","Sinha","Kapoor"]
CITIES = ["Mumbai","Delhi","Bengaluru","Chennai","Hyderabad","Pune","Kolkata","Jaipur","Ahmedabad","Kochi"]

def adult_birthday():
    start = date(1955, 1, 1).toordinal()
    end = date(2005, 12, 31).toordinal()
    return date.fromordinal(random.randint(start, end))

def phone():
    return "+91-" + str(random.randint(70,99)) + str(random.randint(10000000, 99999999))

def height_weight(gender):
    h = random.gauss(172 if gender == "Male" else 160, 7)
    bmi = random.gauss(23.5, 3)
    w = bmi * (h/100)**2
    return round(h,1), round(w,1), round(bmi,1)

rows = []
for i in range(1, N+1):
    gender = random.choice(["Male","Female"])
    first = random.choice(FIRST_NAMES)
    last = random.choice(LAST_NAMES)
    dob = adult_birthday()
    age = (date.today() - dob).days // 365
    h, w, bmi = height_weight(gender)
    rows.append({
        "person_id": f"HM-{i:05d}",
        "name": f"{first} {last}",
        "gender": gender,
        "date_of_birth": dob.isoformat(),
        "age": age,
        "height_cm": h,
        "weight_kg": w,
        "bmi": bmi,
        "blood_group": random.choice(["A+","A-","B+","B-","O+","O-","AB+","AB-"]),
        "city": random.choice(CITIES),
        "phone": phone(),
        "email": f"{first.lower()}.{last.lower()}{random.randint(1,999)}@example.com",
        "occupation": random.choice(["Engineer","Teacher","Doctor","Designer","Analyst","Manager","Artist","Trader","Nurse","Developer"]),
        "salary_inr_lpa": round(random.triangular(3, 45, 10), 1),
        "last_active": (date.today() - timedelta(days=random.randint(0, 90))).isoformat(),
    })

with open("human_models.csv", "w", newline="", encoding="utf-8") as f:
    wtr = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    wtr.writeheader()
    wtr.writerows(rows)
print(f"CSV written: human_models.csv ({len(rows)} records)")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Human_Models"
    headers = list(rows[0].keys())
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="4472C4")
    for r in rows:
        ws.append([r[k] for k in headers])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    ws.freeze_panes = "A2"
    wb.save("human_models.xlsx")
    print("Excel written: human_models.xlsx")
except ImportError:
    print("openpyxl not installed - run: pip install openpyxl --user --break-system-packages")

print(f"Sample record: {rows[0]}")
print("Integrity check: names are synthetic; no real-person data. Rows:", len(rows))
