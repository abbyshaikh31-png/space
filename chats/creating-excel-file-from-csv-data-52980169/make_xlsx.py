import csv, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

DATA = """Name,Home Address,Parent Home Address
Liam Smith,472 Maple Rd Austin TX 78701,109 Cedar Ln Dallas TX 75201
Noah Johnson,815 Oak Ave Denver CO 80202,324 Pine St Boulder CO 80301
 Oliver Williams,234 Pine St Miami FL 33101,657 Elm Dr Orlando FL 32801
Elijah Brown,657 Elm Dr Seattle WA 98101,891 Birch Ct Tacoma WA 98402
James Jones,109 Cedar Ln Chicago IL 60601,472 Maple Rd Naperville IL 60540
William Garcia,891 Birch Ct Phoenix AZ 85001,516 Willow Way Tucson AZ 85701
Benjamin Miller,516 Willow Way Boston MA 02108,743 Walnut Dr Worcester MA 01604
Lucas Davis,743 Walnut Dr Denver CO 80202,234 Pine St Colorado Springs CO 80903
Henry Rodriguez,928 Ash Blvd San Diego CA 92101,815 Oak Ave Los Angeles CA 90001
Alexander Martinez,305 Cherry Ln Atlanta GA 30303,928 Ash Blvd Savannah GA 31401
Mason Hernandez,142 Spruce St New York NY 10001,305 Cherry Ln Buffalo NY 14202
Michael Lopez,689 Hickory Dr Houston TX 77002,142 Spruce St San Antonio TX 78205
Ethan Gonzalez,573 Redwood Ave Portland OR 97201,689 Hickory Dr Eugene OR 97401
Daniel Wilson,412 Magnolia Ct Columbus OH 43215,573 Redwood Ave Cleveland OH 44114
Jacob Anderson,824 Cypress Rd Detroit MI 48226,412 Magnolia Ct Grand Rapids MI 49503
Logan Thomas,256 Poplar Ln Nashville TN 37203,824 Cypress Rd Memphis TN 38103
Jackson Taylor,719 Alder St Charlotte NC 28202,256 Poplar Ln Raleigh NC 27601
Levi Moore,941 Beech Dr Indianapolis IN 46204,719 Alder St Fort Wayne IN 46802
Sebastian Jackson,365 Chestnut Ave Louisville KY 40202,941 Beech Dr Lexington KY 40502
Jack Martin,128 Sycamore Rd Kansas City MO 64106,365 Chestnut Ave St. Louis MO 63101
Owen Lee,582 Linden Blvd Salt Lake City UT 84111,128 Sycamore Rd Provo UT 84601
Theodore Perez,491 Aspen Ct Las Vegas NV 89101,582 Linden Blvd Reno NV 89501
Aiden Thompson,834 Fir Dr Minneapolis MN 55401,491 Aspen Ct Duluth MN 55802
Samuel White,216 Hickory Ln Milwaukee WI 53202,834 Fir Dr Green Bay WI 54301
Joseph Harris,627 Laurel Rd Baltimore MD 21201,216 Hickory Ln Annapolis MD 21401
John Sanchez,749 Myrtle Ave Richmond VA 23219,627 Laurel Rd Norfolk VA 23510
David Clark,315 Redwood St Philadelphia PA 19107,749 Myrtle Ave Pittsburgh PA 15219
Wyatt Ramirez,852 Cedar Blvd San Jose CA 95112,315 Redwood St Sacramento CA 95814
Carter Lewis,964 Maple Ct Oklahoma City OK 73102,852 Cedar Blvd Tulsa OK 74103
Julian Robinson,137 Elm Ave Albuquerque NM 87102,964 Maple Ct Santa Fe NM 87501
Luke Walker,529 Pine Rd Omaha NE 68102,137 Elm Ave Lincoln NE 68508
Burke Young,462 Oak Dr Wichita KS 67202,529 Pine Rd Topeka KS 66603
Santiago Allen,781 Birch Ln Des Moines IA 50309,462 Oak Dr Cedar Rapids IA 52401
Axel King,293 Willow Rd Little Rock AR 72201,781 Birch Ln Fayetteville AR 72701
Angel Wright,614 Walnut Ave Jackson MS 39201,293 Willow Rd Gulfport MS 39501
Roman Scott,847 Ash St Birmingham AL 35203,614 Walnut Ave Mobile AL 36602
Stetson Torres,359 Cherry Blvd Baton Rouge LA 70801,847 Ash St New Orleans LA 70112
Brooks Nguyen,126 Spruce Rd Boise ID 83702,359 Cherry Blvd Idaho Falls ID 83401
Bennett Hill,973 Cypress Ave Cheyenne WY 82001,126 Spruce Rd Casper WY 82601
Jameson Flores,485 Poplar Dr Fargo ND 58102,973 Cypress Ave Bismarck ND 58501
Colton Green,639 Alder Ln Sioux Falls SD 57104,485 Poplar Dr Rapid City SD 57701
Ian Garcia,152 Beech St Billings MT 59101,639 Alder Ln Missoula MT 59801
Clara Martinez,794 Chestnut Rd Manchester NH 03101,152 Beech St Concord NH 03301
Eleanor Robinson,825 Sycamore Ave Burlington VT 05401,794 Chestnut Rd Montpelier VT 05602
Lucy Clark,341 Linden Dr Portland ME 04101,825 Sycamore Ave Augusta ME 04330
Cora Rodriguez,916 Aspen Rd Providence RI 02903,341 Linden Dr Newport RI 02840
Eva Carter,562 Fir Ave Wilmington DE 19801,916 Aspen Rd Dover DE 19901
Alice Mitchell,274 Hickory Blvd Charleston WV 25301,562 Fir Ave Morgantown WV 26505
Amelia Roberts,893 Laurel Ct Anchorage AK 99501,274 Hickory Blvd Fairbanks AK 99701
Mateo Turner,415 Myrtle Rd Honolulu HI 96813,893 Laurel Ct Hilo HI 96720"""

rows = list(csv.reader(io.StringIO(DATA)))
rows = [[c.strip() for c in r] for r in rows if any(c.strip() for c in r)]

wb = Workbook()
ws = wb.active
ws.title = "Addresses"
for r in rows:
    ws.append(r)

hdr_fill = PatternFill("solid", fgColor="2C3E50")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = hdr_fill
ws.freeze_panes = "A2"
for col in ws.columns:
    letter = col[0].column_letter
    ws.column_dimensions[letter].width = max(len(str(c.value or "")) for c in col) + 3

wb.save("addresses.xlsx")
print("OK rows:", len(rows) - 1)
